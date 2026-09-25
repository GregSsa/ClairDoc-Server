import csv
import html
import re
from email import policy
from email.parser import BytesParser
from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation

TEXT_EXTENSIONS = {".txt", ".md", ".csv", ".tsv", ".log"}
OFFICE_EXTENSIONS = {".docx", ".xlsx", ".pptx"}
EMAIL_EXTENSIONS = {".eml"}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp"}
SUPPORTED_EXTENSIONS = (
    {".pdf"} | TEXT_EXTENSIONS | OFFICE_EXTENSIONS | EMAIL_EXTENSIONS | IMAGE_EXTENSIONS
)


def extract_document_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix in {".txt", ".md", ".log"}:
        return path.read_text(encoding="utf-8", errors="replace")
    if suffix in {".csv", ".tsv"}:
        delimiter = "\t" if suffix == ".tsv" else ","
        with path.open(encoding="utf-8", errors="replace", newline="") as source:
            return "\n".join(" | ".join(row) for row in csv.reader(source, delimiter=delimiter))
    if suffix == ".docx":
        document = Document(path)
        return "\n".join(paragraph.text for paragraph in document.paragraphs)
    if suffix == ".xlsx":
        workbook = load_workbook(path, read_only=True, data_only=True)
        parts: list[str] = []
        for sheet in workbook.worksheets:
            parts.append(f"Feuille : {sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                parts.append(" | ".join("" if value is None else str(value) for value in row))
        workbook.close()
        return "\n".join(parts)
    if suffix == ".pptx":
        presentation = Presentation(path)
        return "\n".join(
            shape.text
            for slide in presentation.slides
            for shape in slide.shapes
            if hasattr(shape, "text") and shape.text
        )
    if suffix == ".eml":
        message = BytesParser(policy=policy.default).parsebytes(path.read_bytes())
        parts = [f"Objet : {message.get('subject', '')}", f"De : {message.get('from', '')}"]
        bodies = message.walk() if message.is_multipart() else [message]
        for part in bodies:
            if part.get_content_type() not in {"text/plain", "text/html"}:
                continue
            body = part.get_content()
            if part.get_content_type() == "text/html":
                body = re.sub(r"<[^>]+>", " ", html.unescape(body))
            parts.append(body)
        return "\n".join(parts)
    raise ValueError(f"Format non pris en charge : {suffix}")
